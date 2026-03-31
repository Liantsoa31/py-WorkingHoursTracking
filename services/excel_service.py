from openpyxl import Workbook, load_workbook
from config.settings import EXCEL_PATH, TABLE_HEADERS
from services.date_service import get_month_key, get_month_label

class ExcelService:

    @staticmethod
    def save(date, name, number):
        if EXCEL_PATH.exists():
            wb = load_workbook(EXCEL_PATH)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(TABLE_HEADERS)

        ws.append([date, name, int(number)])
        wb.save(EXCEL_PATH)

    @staticmethod
    def get_all():
        if not EXCEL_PATH.exists():
            return []

        wb = load_workbook(EXCEL_PATH)
        ws = wb.active

        data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            data.append(row)

        return data

    @staticmethod
    def get_months():
        data = ExcelService.get_all()

        months = {}
        for row in data:
            date_str = row[0]
            key = get_month_key(date_str)
            label = get_month_label(date_str)
            months[key] = label  # évite les doublons

        return months  # dict { "2026-01": "Janvier 2026" }

    @staticmethod
    def delete_row(row_index):
        """
        Supprime une ligne spécifique dans le fichier Excel.
        row_index : ligne dans Excel (1 = header)
        """
        if not EXCEL_PATH.exists():
            return

        wb = load_workbook(EXCEL_PATH)
        ws = wb.active

        # ligne 1 = header, ligne_index >= 2
        if row_index < 2 or row_index > ws.max_row:
            return

        ws.delete_rows(row_index)
        wb.save(EXCEL_PATH)

    @staticmethod
    def delete_all():
        if not EXCEL_PATH.exists():
            return

        wb = load_workbook(EXCEL_PATH)
        ws = wb.active

        ws.delete_rows(2, ws.max_row)  # garde la ligne d’en-tête
        wb.save(EXCEL_PATH)